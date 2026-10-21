"""Extract text from an uploaded document, page by page (FR-14).

Uploaded files are untrusted data (NFR-07). Nothing here executes or interprets
their content: it is parsed into plain text and nothing more. A file that cannot
be parsed raises a handled `ExtractionError`; it never reaches the database.
"""

import hashlib
import io
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from docx import Document
from pypdf import PdfReader

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".md", ".txt"})
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
PAGE_SEPARATOR = "\f"


class UnsupportedFormatError(Exception):
    """The file's extension is not one of `SUPPORTED_SUFFIXES`."""


class ExtractionError(Exception):
    """The file is corrupt, unreadable, or holds no text."""


@dataclass(frozen=True)
class ExtractedDocument:
    """The text of a document, one entry per page.

    Attributes:
        pages: Page texts in order. Formats without pages (DOCX, Markdown, plain
            text) have exactly one.
    """

    pages: tuple[str, ...]

    @property
    def text(self) -> str:
        """All pages joined by a form feed, the layout `store_document` writes to disk."""
        return PAGE_SEPARATOR.join(self.pages)


def suffix_of(filename: str) -> str:
    """Return the lower-cased extension of `filename`, ignoring any directory part."""
    return PurePosixPath(filename.replace("\\", "/")).suffix.lower()


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    """Extract the text of an uploaded document.

    Args:
        filename: The uploaded name; only its extension is used, and it is checked
            before `data` is touched.
        data: The file's bytes.

    Returns:
        The extracted pages.

    Raises:
        UnsupportedFormatError: If the extension is not supported.
        ExtractionError: If the file cannot be parsed or holds no text.
    """
    suffix = suffix_of(filename)
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise UnsupportedFormatError(
            f"{suffix or 'a file without an extension'} is not supported; use one of {supported}"
        )
    try:
        pages = _READERS[suffix](data)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"could not read {suffix} file: {exc}") from exc
    cleaned = tuple(page.replace(PAGE_SEPARATOR, " ").strip() for page in pages)
    if not any(cleaned):
        raise ExtractionError("the document has no extractable text")
    return ExtractedDocument(cleaned)


def _read_pdf(data: bytes) -> list[str]:
    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]


def _read_docx(data: bytes) -> list[str]:
    return ["\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)]


def _read_text(data: bytes) -> list[str]:
    return [data.decode("utf-8")]


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".md": _read_text, ".txt": _read_text}


@dataclass(frozen=True)
class StoredDocument:
    """Where an uploaded document's bytes and extracted text were written.

    Attributes:
        sha256: Hex digest of the original bytes; also the stem of both files.
        original: Path of the stored original upload.
        text: Path of the extracted text, pages separated by `PAGE_SEPARATOR`.
    """

    sha256: str
    original: Path
    text: Path


def store_document(
    data_dir: Path, filename: str, data: bytes, document: ExtractedDocument
) -> StoredDocument:
    """Write the original bytes and the extracted text under `data_dir/documents/`.

    Files are named by the SHA-256 of the bytes, never by the uploaded name, so a
    hostile filename cannot pick where anything is written, and re-uploading the
    same file is idempotent.

    Args:
        data_dir: The private data directory (gitignored).
        filename: The uploaded name; only its extension is kept.
        data: The original bytes.
        document: The text extracted from `data`.

    Returns:
        The paths written.
    """
    digest = hashlib.sha256(data).hexdigest()
    folder = data_dir / "documents"
    folder.mkdir(parents=True, exist_ok=True)
    original = folder / f"{digest}{suffix_of(filename)}"
    text = folder / f"{digest}.pages.txt"
    original.write_bytes(data)
    text.write_text(document.text, encoding="utf-8")
    return StoredDocument(digest, original, text)
